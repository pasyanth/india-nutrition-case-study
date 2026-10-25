"""Build the formatted assignment DOCX from the final Markdown report."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
MD_PATH = ROOT / "final_case_study.md"
OUT_PATH = ROOT / "India_Nutrition_Protein_Case_Study.docx"


def move_references_to_end(text):
    supplement = re.search(r"(?m)^## Supplemental protein analyses\s*$", text)
    refs = re.search(r"(?m)^## References\s*$", text)
    if not refs or not supplement:
        raise ValueError("Could not find the References and Supplemental protein analyses sections")
    if refs.start() > supplement.start():
        return text
    block_end = re.search(r"(?m)^## ", text[refs.end():])
    end = refs.end() + block_end.start() if block_end else len(text)
    refs_block = text[refs.start():end].rstrip()
    text = text[:refs.start()] + text[end:]
    return text.rstrip() + "\n\n" + refs_block + "\n"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "5")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9E0E7")


def add_hyperlink(paragraph, label, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "176B87"); props.append(color)
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single"); props.append(underline)
    run.append(props)
    text = OxmlElement("w:t"); text.text = label; run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text):
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2]); run.bold = True
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1]); run.italic = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1]); run.font.name = "Consolas"; run.font.size = Pt(9)
        else:
            link = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            add_hyperlink(paragraph, link.group(1), link.group(2))
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def is_table_row(line):
    return line.strip().startswith("|") and line.strip().endswith("|")


def split_table_row(line):
    return [x.strip() for x in line.strip().strip("|").split("|")]


def add_markdown_table(doc, lines):
    parsed = [split_table_row(line) for line in lines]
    body = [row for i, row in enumerate(parsed) if i != 1]
    cols = max(len(row) for row in body)
    table = doc.add_table(rows=len(body), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    set_table_borders(table)
    for ri, row in enumerate(body):
        for ci in range(cols):
            cell = table.cell(ri, ci)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if ci < len(row):
                add_inline(p, row[ci])
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.name = "Arial"
            if ri == 0:
                set_cell_shading(cell, "173B57")
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.size = Pt(9)
            elif ri % 2 == 0:
                set_cell_shading(cell, "F2F6F9")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_image(doc, alt, relpath):
    image_path = (ROOT / relpath).resolve()
    if not image_path.exists():
        raise FileNotFoundError(f"Missing image: {image_path}")
    from PIL import Image
    with Image.open(image_path) as im:
        ratio = im.height / im.width
    width = 4.65 if ratio > 1.0 else 6.35
    if ratio > 1.0:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    shape = run.add_picture(str(image_path), width=Inches(width))
    shape._inline.docPr.set("descr", alt)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(10)
    r = cap.add_run(alt); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(90, 98, 105)


def build_docx(text):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(35, 42, 48)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in (("Title", 21), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        st = styles[name]
        st.font.name = "Arial"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(14 if name != "Title" else 0)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)

    lines = text.splitlines()
    i = 0
    title_seen = False
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1; continue
        if line.startswith("!["):
            m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
            if m:
                add_image(doc, m.group(1), m.group(2))
                i += 1; continue
        if is_table_row(line):
            block = []
            while i < len(lines) and is_table_row(lines[i].strip()):
                block.append(lines[i].strip()); i += 1
            add_markdown_table(doc, block)
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            label = heading.group(2).strip()
            if level == 1 and not title_seen:
                p = doc.add_paragraph(style="Title")
                add_inline(p, label)
                title_seen = True
            else:
                p = doc.add_paragraph(style=f"Heading {level}")
                add_inline(p, label)
            i += 1; continue
        if re.match(r"^[-*]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, re.sub(r"^[-*]\s+", "", line))
            p.paragraph_format.space_after = Pt(3)
            i += 1; continue
        if re.match(r"^\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\.\s+", "", line))
            p.paragraph_format.space_after = Pt(3)
            i += 1; continue
        p = doc.add_paragraph()
        add_inline(p, line)
        i += 1

    doc.core_properties.title = "Indian Nutrition Trends Protein Case Study"
    doc.core_properties.subject = "Protein patterns in India's HCES 2022-23 and 2023-24 estimates"
    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    source = MD_PATH.read_text(encoding="utf-8")
    arranged = move_references_to_end(source)
    MD_PATH.write_text(arranged, encoding="utf-8")
    print(build_docx(arranged))
