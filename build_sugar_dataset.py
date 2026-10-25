"""Build HCES household-level sugar/jaggery quantities and weighted summaries.

Raw survey CSVs are read only. Food quantities are kilograms for the last 30
days. This measures sugar commodities, not total sugars from all foods.
"""

from __future__ import annotations

import csv
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "outputs"
ID_POSITIONS = tuple(range(2, 16))  # FSU through sample household identifiers
TARGET_CODES = {"074", "171", "172", "173", "174", "175"}
LABELS = {
    "074": "sugar_free_pds", "171": "sugar_pds_subsidized",
    "172": "sugar_other_sources", "173": "jaggery_gur",
    "174": "candy_misri", "175": "honey",
}


def find_csv(folder: str, prefix: str) -> Path:
    matches = [p for p in (ROOT / folder).rglob("*.csv") if p.name.lower().startswith(prefix.lower())]
    if len(matches) != 1:
        raise FileNotFoundError(f"Expected one {prefix} CSV in {folder}, found {len(matches)}")
    return matches[0]


def key_from_row(row: list[str]) -> tuple[str, ...]:
    parts = []
    for pos in ID_POSITIONS:
        value = row[pos].strip()
        # The two delivered CSV conversions differ on zero padding for numeric codes.
        if value.endswith(".0") and value[:-2].isdigit():
            value = value[:-2]
        if value.isdigit():
            value = value.lstrip("0") or "0"
        parts.append(value)
    return tuple(parts)


def header_index(headers: list[str], *names: str) -> int:
    by_lower = {h.lower(): i for i, h in enumerate(headers)}
    for name in names:
        if name.lower() in by_lower:
            return by_lower[name.lower()]
    raise KeyError(f"Could not find one of {names} in {headers}")


def process_year(folder: str, year: str) -> tuple[int, list[dict[str, str]]]:
    household_file = find_csv(folder, "LEVEL - 03")
    food_file = find_csv(folder, "LEVEL - 05")
    state_ws = load_workbook(ROOT / "tabulation_state_code_2.xlsx", read_only=True, data_only=True).active
    state_names = {str(row[0]).zfill(2): str(row[1]) for row in state_ws.iter_rows(min_row=2, values_only=True) if row[0] is not None}
    households: dict[tuple[str, ...], tuple[str, str, str]] = {}

    with household_file.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream)
        headers = next(reader)
        size_idx = header_index(headers, "HH_Size_FDQ", "b2q2pt1")
        for row in reader:
            households[key_from_row(row)] = (row[4].strip().zfill(2), row[3].strip(), row[size_idx].strip())

    food_weights: dict[tuple[str, ...], float] = {}
    quantities: dict[tuple[str, ...], dict[str, float]] = defaultdict(lambda: defaultdict(float))
    invalid_qty = 0
    with food_file.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream)
        headers = next(reader)
        item_idx = header_index(headers, "Item_Code", "b5pt1q2")
        # 2022-23 uses positional questionnaire field names: q5 is total
        # consumption quantity; q7 is the item-source code. 2023-24 is labeled.
        qty_idx = header_index(headers, "Total_Consumption_Quantity", "b5pt1q5")
        weight_idx = header_index(headers, "Multiplier", "mult")
        for row in reader:
            key = key_from_row(row)
            if key not in food_weights:
                try:
                    food_weights[key] = float(row[weight_idx])
                except ValueError:
                    continue
            code = row[item_idx].strip().zfill(3)
            if code not in TARGET_CODES:
                continue
            raw_qty = row[qty_idx].strip()
            if not raw_qty:
                continue
            try:
                amount = float(raw_qty)
                if amount < 0:
                    invalid_qty += 1
                    continue
                quantities[key][code] += amount
            except ValueError:
                invalid_qty += 1

    unmatched_keys = set(food_weights).difference(households)
    missing_food_rows = set(households).difference(food_weights)
    excluded = 0
    domains: dict[tuple[str, str], dict[str, float]] = defaultdict(lambda: defaultdict(float))
    household_path = OUT / f"sugar_household_level_{year}.csv"
    columns = [
        "year", "state_code", "state_ut", "sector_code", "fsu", "nss_region", "district",
        "stratum", "sub_stratum", "panel", "sub_sample", "fod_subregion", "sample_su",
        "sample_subdivision", "second_stage_stratum", "sample_household",
        "household_size_fdq", "fdq_multiplier",
        "sugar_kg_30d", "sugar_g_per_person_per_day", "jaggery_gur_kg_30d",
        "candy_misri_kg_30d", "honey_kg_30d",
    ] + [LABELS[code] + "_kg_30d" for code in ("074", "171", "172")]
    matched = 0
    with household_path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns)
        writer.writeheader()
        for key, weight in food_weights.items():
            hh = households.get(key)
            if hh is None:
                continue
            state, sector, raw_size = hh
            try:
                size = float(raw_size)
                if size <= 0 or weight <= 0:
                    raise ValueError
            except ValueError:
                excluded += 1
                continue
            amounts = {code: quantities.get(key, {}).get(code, 0.0) for code in TARGET_CODES}
            sugar = amounts["074"] + amounts["171"] + amounts["172"]
            record = {
                "year": year, "state_code": state, "state_ut": state_names.get(state, "Unknown"), "sector_code": sector,
                "fsu": key[0], "nss_region": key[3], "district": key[4],
                "stratum": key[5], "sub_stratum": key[6], "panel": key[7],
                "sub_sample": key[8], "fod_subregion": key[9], "sample_su": key[10],
                "sample_subdivision": key[11], "second_stage_stratum": key[12],
                "sample_household": key[13],
                "household_size_fdq": f"{size:g}", "fdq_multiplier": f"{weight:g}",
                "sugar_kg_30d": f"{sugar:.6f}",
                "sugar_g_per_person_per_day": f"{sugar * 1000 / 30 / size:.6f}",
                "jaggery_gur_kg_30d": f"{amounts['173']:.6f}",
                "candy_misri_kg_30d": f"{amounts['174']:.6f}",
                "honey_kg_30d": f"{amounts['175']:.6f}",
            }
            for code, label in LABELS.items():
                record[label + "_kg_30d"] = f"{amounts[code]:.6f}"
            writer.writerow(record)
            matched += 1
            for domain in [("All-India", "All"), ("All-India", sector), (state, sector)]:
                agg = domains[domain]
                agg["weighted_quantity"] += weight * sugar
                agg["weighted_population"] += weight * size
                agg["weighted_households"] += weight
                agg["n"] += 1

    summaries = []
    for (state, sector), agg in domains.items():
        summaries.append({
            "year": year, "state_code": state,
            "state_ut": "All-India" if state == "All-India" else state_names.get(state, "Unknown"),
            "sector_code": sector,
            "households_unweighted": str(int(agg["n"])),
            "weighted_sugar_g_per_person_per_day": f"{agg['weighted_quantity'] * 1000 / 30 / agg['weighted_population']:.6f}",
            "weighted_sugar_kg_per_household_per_30d": f"{agg['weighted_quantity'] / agg['weighted_households']:.6f}",
        })
    print(
        f"{year}: food keys={len(food_weights):,}; matched households={matched:,}; "
        f"unmatched food keys={len(unmatched_keys):,}; households without food rows={len(missing_food_rows):,}; "
        f"excluded size/weight={excluded:,}; invalid target quantities={invalid_qty:,}", flush=True
    )
    return matched, summaries


def main() -> None:
    OUT.mkdir(exist_ok=True)
    result_22 = process_year("CSV_data_HH_Cons_exp_22_23", "2022-23")
    result_23 = process_year("HCES_Data_2023-24_Csv", "2023-24")
    rows = result_22[1] + result_23[1]
    fields = [
        "year", "state_code", "state_ut", "sector_code", "households_unweighted",
        "weighted_sugar_g_per_person_per_day", "weighted_sugar_kg_per_household_per_30d",
    ]
    with (OUT / "sugar_weighted_summary.csv").open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
    print(f"Wrote {result_22[0] + result_23[0]:,} household rows and {len(rows):,} summary rows.", flush=True)


if __name__ == "__main__":
    main()
